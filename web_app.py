# -*- coding: utf-8 -*-
"""
Web App: Phân bổ chi phí điện - GreenNode AgentBase
API endpoints for processing PDF/Excel/Image and updating Excel
"""
import os
import re
import io
import tempfile
import shutil
from datetime import datetime
import calendar

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import JSONResponse, FileResponse
from pydantic import BaseModel

import pdfplumber
import openpyxl

# Optional imports
try:
    import pytesseract
    from PIL import Image
    HAS_OCR = True
except ImportError:
    HAS_OCR = False

# ==========================================
# CONFIG
# ==========================================
DATA_START_ROW = 5
COL_CODE = 8
COL_KY_TT = 13
COL_KY_TT_TRUOC = 14
COL_TONG_TIEN_DIEM = 15
EXCEL_PREFIX = "Bảng phân bổ thanh toán chi phí điện"
DEFAULT_DIR = "/data"

# ==========================================
# APP
# ==========================================
app = FastAPI(title="Phân bổ chi phí điện", version="1.0")

@app.get("/health")
async def health():
    return {"status": "ok"}

@app.get("/")
async def root():
    return {"app": "Phân bổ chi phí điện", "version": "1.0", "health": "/health", "docs": "/docs"}

# ==========================================
# LOGIC
# ==========================================

def add_months(date_str, n):
    try:
        d = datetime.strptime(date_str.strip(), '%d/%m/%Y')
        total_months = d.year * 12 + d.month - 1 + n
        year = total_months // 12
        month = total_months % 12 + 1
        last_day = calendar.monthrange(year, month)[1]
        day = min(d.day, last_day)
        return d.replace(year=year, month=month, day=day).strftime('%d/%m/%Y')
    except:
        return None

def shift_date_range(date_range_str, n):
    if not date_range_str:
        return None
    match = re.match(r'(\d{2}/\d{2}/\d{4})\s*-\s*(\d{2}/\d{2}/\d{4})', str(date_range_str).strip())
    if not match:
        return None
    s = add_months(match.group(1), n)
    e = add_months(match.group(2), n)
    if s and e:
        return f"{s} - {e}"
    return None

def _extract_code(text):
    if not text:
        return ""
    text = str(text).strip()
    m = re.match(r'^([A-Z]{2}[0-9A-Z]+)', text)
    if not m:
        m = re.search(r'\b([A-Z]{2}\d{8,})\b', text)
    return m.group(1) if m else ""

def _extract_amount(debit, credit):
    s = ""
    if debit and debit != '0':
        s = debit.replace(',', '').replace('.00', '').strip()
    elif credit and credit != '0':
        s = credit.replace(',', '').replace('.00', '').strip()
    return float(s) if s else 0

def extract_pdf_data(filepath):
    data = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                for row in table:
                    if not row or not row[0]:
                        continue
                    stt = str(row[0]).strip() if row[0] else ""
                    if not stt.isdigit():
                        continue
                    desc = str(row[5]).strip() if row[5] else ""
                    debit = str(row[6]).strip() if row[6] else ""
                    credit = str(row[7]).strip() if row[7] else ""
                    code = _extract_code(desc)
                    amount = _extract_amount(debit, credit)
                    if code and code != 'MSB' and amount > 0:
                        data.append({'code': code, 'amount': amount})
    return data

def extract_excel_data(filepath):
    data = []
    wb = openpyxl.load_workbook(filepath, keep_vba=False, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=5, values_only=True):
            if not row or not row[0]:
                continue
            code = str(row[0]).strip()
            amount = 0
            if len(row) > 2 and row[2]:
                try:
                    amount = float(str(row[2]).replace(',', '').replace('.00', '').strip())
                except:
                    pass
            if code and amount > 0:
                data.append({'code': code, 'amount': amount})
    wb.close()
    return data

def extract_image_data(filepath):
    if not HAS_OCR:
        raise HTTPException(status_code=500, detail="OCR not available. Install pytesseract.")
    tess_path = os.environ.get('TESSERACT_CMD', '/usr/bin/tesseract')
    pytesseract.pytesseract.tesseract_cmd = tess_path
    img = Image.open(filepath)
    text = pytesseract.image_to_string(img, lang='eng')
    data = []
    for line in text.split('\n'):
        code = _extract_code(line)
        nums = re.findall(r'[\d,]+\.?\d*', line)
        amount = float(nums[-1].replace(',', '')) if nums else 0
        if code and code != 'MSB' and amount > 0:
            data.append({'code': code, 'amount': amount})
    return data

def build_lookup(data_list):
    lookup = {}
    for item in data_list:
        if item['code'] not in lookup:
            lookup[item['code']] = item
    return lookup

def process_excel(excel_path, lookup, month_shift=0):
    wb = openpyxl.load_workbook(excel_path, keep_vba=False)
    ws = wb.worksheets[1]
    results = []
    updated = skipped_dd = skipped_nm = 0

    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        code = ws.cell(row=row_idx, column=COL_CODE).value
        if not code:
            continue
        code = str(code).strip()
        c13 = ws.cell(row=row_idx, column=COL_KY_TT).value or ""
        c14 = ws.cell(row=row_idx, column=COL_KY_TT_TRUOC).value or ""

        if any(kw in str(c13).lower() for kw in ['di d', 'dời', 'ngừng']) or \
           any(kw in str(c14).lower() for kw in ['di d', 'dời', 'ngừng']):
            ws.cell(row=row_idx, column=COL_TONG_TIEN_DIEM).value = None
            skipped_dd += 1
            results.append({'row': row_idx, 'code': code, 'action': 'SKIP', 'amount': 'blank'})
            continue

        if month_shift != 0:
            n13 = shift_date_range(c13, month_shift)
            if n13:
                ws.cell(row=row_idx, column=COL_KY_TT).value = n13
            n14 = shift_date_range(c14, month_shift)
            if n14:
                ws.cell(row=row_idx, column=COL_KY_TT_TRUOC).value = n14

        if code in lookup:
            ws.cell(row=row_idx, column=COL_TONG_TIEN_DIEM).value = lookup[code]['amount']
            updated += 1
            results.append({'row': row_idx, 'code': code, 'action': 'UPDATED', 'amount': lookup[code]['amount']})
        else:
            skipped_nm += 1
            results.append({'row': row_idx, 'code': code, 'action': 'NO_MATCH', 'amount': 0})

    return wb, results, updated, skipped_dd, skipped_nm

# ==========================================
# API ENDPOINTS
# ==========================================

class ProcessResponse(BaseModel):
    updated: int
    skipped_di_doi: int
    no_match: int
    total: int
    results: list

@app.post("/process", response_model=ProcessResponse)
async def process(
    input_file: UploadFile = File(...),
    excel_template: UploadFile = File(...),
    month: int = Form(...),
):
    """Process input file (PDF/Excel/Image) and update Excel template."""
    # Save uploaded files to temp
    tmpdir = tempfile.mkdtemp()
    try:
        input_path = os.path.join(tmpdir, input_file.filename)
        with open(input_path, 'wb') as f:
            f.write(await input_file.read())

        excel_path = os.path.join(tmpdir, excel_template.filename)
        with open(excel_path, 'wb') as f:
            f.write(await excel_template.read())

        # Extract data based on file type
        ext = os.path.splitext(input_file.filename)[1].lower()
        if ext == '.pdf':
            data = extract_pdf_data(input_path)
        elif ext in ('.xlsx', '.xls'):
            data = extract_excel_data(input_path)
        elif ext in ('.jpg', '.jpeg', '.png', '.gif'):
            data = extract_image_data(input_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}")

        lookup = build_lookup(data)
        month_shift = month - 7
        wb, results, updated, skipped_dd, skipped_nm = process_excel(excel_path, lookup, month_shift)

        # Update title
        month_num = str(month).zfill(2)
        ws = wb.worksheets[1]
        ws.cell(row=3, column=1).value = f"BẢNG PHÂN BỔ CHI PHÍ THÁNG {month_num}/2026"
        ws.title = f"Phân bổ_{month_num}.26"

        # Save output
        output_name = f"{EXCEL_PREFIX} tháng {month}.2026.xlsx"
        output_path = os.path.join(tmpdir, output_name)
        wb.save(output_path)
        wb.close()

        return ProcessResponse(
            updated=updated,
            skipped_di_doi=skipped_dd,
            no_match=skipped_nm,
            total=updated + skipped_dd + skipped_nm,
            results=results[:50],
        )
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

@app.post("/process/download")
async def process_download(
    input_file: UploadFile = File(...),
    excel_template: UploadFile = File(...),
    month: int = Form(...),
):
    """Process and return the updated Excel file for download."""
    tmpdir = tempfile.mkdtemp()
    try:
        input_path = os.path.join(tmpdir, input_file.filename)
        with open(input_path, 'wb') as f:
            f.write(await input_file.read())

        excel_path = os.path.join(tmpdir, excel_template.filename)
        with open(excel_path, 'wb') as f:
            f.write(await excel_template.read())

        ext = os.path.splitext(input_file.filename)[1].lower()
        if ext == '.pdf':
            data = extract_pdf_data(input_path)
        elif ext in ('.xlsx', '.xls'):
            data = extract_excel_data(input_path)
        elif ext in ('.jpg', '.jpeg', '.png', '.gif'):
            data = extract_image_data(input_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported: {ext}")

        lookup = build_lookup(data)
        month_shift = month - 7
        wb, results, updated, skipped_dd, skipped_nm = process_excel(excel_path, lookup, month_shift)

        month_num = str(month).zfill(2)
        ws = wb.worksheets[1]
        ws.cell(row=3, column=1).value = f"BẢNG PHÂN BỔ CHI PHÍ THÁNG {month_num}/2026"
        ws.title = f"Phân bổ_{month_num}.26"

        output_name = f"{EXCEL_PREFIX} tháng {month}.2026.xlsx"
        output_path = os.path.join(tmpdir, output_name)
        wb.save(output_path)
        wb.close()

        return FileResponse(output_path, filename=output_name,
                            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8080)
